"""
Excel ต้องไม่ให้เป้าเงิน 0 กับพนักงานที่แยกหลายคลัง (ผลตรวจ 28 ก.ย. 2026 §3.2)

หน้าเว็บเก็บเป้าเงินไว้ที่คีย์ _allocKey ซึ่งเป็น "รหัส|คลัง" (C442|R408) สำหรับคนที่แยกคลัง
เดิมส่งคีย์นั้นไปเป็น emp_id ตรง ๆ ฝั่ง server หาด้วยรหัสพนักงานอย่างเดียว → เป้าเงินรายคน 0
ทั้งที่ยอดรวมหัวตารางถูก — ตัวเลขสองจุดในไฟล์เดียวขัดกันเอง

ทำงานในโฟลเดอร์ชั่วคราว ไม่ยิงเน็ต
"""

from __future__ import annotations

import os
import re
import shutil
import sys
import tempfile
import unittest

import openpyxl
import pandas as pd

REPO = os.path.normpath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, REPO)

from backend.generate_excel import GREEN_FILL, create_target_excel  # noqa: E402


class TestYellowForMultiWarehouseEmployee(unittest.TestCase):
    def setUp(self):
        self._cwd = os.getcwd()
        self._tmpdir = tempfile.mkdtemp(prefix="yellow_wh_")
        os.makedirs(os.path.join(self._tmpdir, "data"), exist_ok=True)
        os.chdir(self._tmpdir)

    def tearDown(self):
        os.chdir(self._cwd)
        shutil.rmtree(self._tmpdir, ignore_errors=True)

    def _write(self):
        rows = []
        for emp, wh, boxes in (("C442", "R408", 3), ("C442", "R409", 2), ("C500", "", 5)):
            rows.append({
                "emp_id": emp, "warehouse_code": wh, "sku": "734046", "allocated_boxes": boxes,
                "hist_avg": 1.0, "hist_ly_same_month": 0.0, "hist_prev_month": 0.0,
                "price_per_box": 1000.0, "brand_name_thai": "ปรุงทิพย์",
                "brand_name_english": "", "product_name_thai": "", "product_name_english": "",
            })
        pd.DataFrame(rows).to_csv("data/r.csv", index=False)
        pd.DataFrame([{"sku": "734046", "price_per_box": 1000.0, "supervisor_target_boxes": 10,
                       "brand_name_thai": "ปรุงทิพย์"}]).to_csv("data/t.csv", index=False)

    def test_money_target_of_split_employee_is_the_sum_of_both_warehouses(self):
        self._write()
        # สิ่งที่หน้าเว็บส่งหลังแก้: ตัด "|คลัง" แล้ว server รวมต่อรหัส (exporting.py)
        sent = [("C442", 3000.0), ("C442", 2000.0), ("C500", 5000.0)]
        yellow = {}
        for emp, v in sent:
            yellow[emp] = yellow.get(emp, 0.0) + v
        out = create_target_excel(
            result_csv="data/r.csv", output_path="data/out.xlsx", brand_filter="ALL",
            yellow_map=yellow, sup_id="SLTEST", target_boxes_csv="data/t.csv",
        )
        self.assertEqual(self._total_fill(out, "C442"), GREEN_FILL.fgColor.rgb,
                         "เป้าเงิน 3,000+2,000 = มูลค่าหีบ 5,000 ต้องขึ้นสีเขียว (ตรงเป้า)")

    def test_old_keys_would_have_shown_a_false_warning(self):
        """พิสูจน์ว่าเทสต์ข้างบนจับบั๊กเดิมได้ — คีย์ "รหัส|คลัง" ทำให้เป้าเงินเป็น 0"""
        self._write()
        out = create_target_excel(
            result_csv="data/r.csv", output_path="data/out.xlsx", brand_filter="ALL",
            yellow_map={"C442|R408": 3000.0, "C442|R409": 2000.0, "C500": 5000.0},
            sup_id="SLTEST", target_boxes_csv="data/t.csv",
        )
        self.assertNotEqual(self._total_fill(out, "C442"), GREEN_FILL.fgColor.rgb)

    @staticmethod
    def _total_fill(path, emp):
        ws = openpyxl.load_workbook(path).active
        for row in ws.iter_rows():
            if len(row) > 1 and str(row[1].value or "").strip() == emp:
                last = max(c.column for c in row if c.value is not None)
                return ws.cell(row=row[0].row, column=last).fill.fgColor.rgb
        raise AssertionError(f"ไม่พบแถวของ {emp}")


class TestFrontendStripsWarehouseFromKey(unittest.TestCase):
    def test_export_payload_sends_plain_employee_codes(self):
        with open(os.path.join(REPO, "frontend", "app.js"), encoding="utf-8") as fh:
            src = fh.read()
        m = re.search(r"yellow_targets: Object\.entries\(S\.yellow\)\.map\(\(\[key, v\]\) => \(\{(.*?)\}\)\)", src, re.S)
        self.assertIsNotNone(m, "ไม่พบจุดประกอบ yellow_targets")
        self.assertIn('String(key).split("|")[0]', m.group(1))


if __name__ == "__main__":
    unittest.main()
